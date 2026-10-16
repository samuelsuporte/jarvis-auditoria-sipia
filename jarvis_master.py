"""Jarvis Master: agente portátil otimizado para auditorias de TI e conformidade SIPIA/Selo UNICEF."""
from __future__ import annotations

import argparse
import datetime as dt
import json
import logging
import os
import platform
import socket
import subprocess
import sys
import threading
import time
import webbrowser
from pathlib import Path
from typing import Any

try:
    import requests
except Exception:
    requests = None
try:
    import psutil
except Exception:
    psutil = None
try:
    from flask import Flask, jsonify, request, send_from_directory
except Exception:
    Flask = None
try:
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).resolve().parent / ".env")
except Exception:
    pass

ROOT = Path(__file__).resolve().parent
LOG_DIR, REPORT_DIR, CONFIG_DIR, KNOWLEDGE_DIR = ROOT / "logs", ROOT / "reports", ROOT / "config", ROOT / "docs"
for folder in (LOG_DIR, REPORT_DIR, CONFIG_DIR, KNOWLEDGE_DIR):
    try:
        folder.mkdir(parents=True, exist_ok=True)
    except Exception:
        pass

logging.basicConfig(filename=str(LOG_DIR / "jarvis_activity.log"), level=logging.INFO,
                    format="%(asctime)s %(levelname)s %(message)s", encoding="utf-8")
log = logging.getLogger("jarvis")

SIPIA_URL = "https://sipiact.mdh.gov.br/"
PERSONA = (
    "Você é um Analista Sênior de Suporte de TI da Secretaria Municipal de Assistência Social de Várzea Grande. "
    "Seu foco atual é atender com urgência a C.I. Nº 277/2026 e o Procedimento SIMP nº 016539-006/2025 (Ministério Público / Selo UNICEF), "
    "avaliando infraestrutura, conectividade, adequação de hardware e acesso ao sistema SIPIA Conselho Tutelar. "
    "Responda em português brasileiro com rigor técnico, clareza administrativa e conformidade legal."
)


def safe(call, fallback=None, label="operação"):
    try:
        return call()
    except Exception as exc:
        log.exception("Falha em %s: %s", label, exc)
        return fallback


def command(args: list[str], timeout: int = 8) -> str:
    def run() -> str:
        # Adicionado encoding='utf-8' e errors='replace' para evitar erros de codec no Windows
        completed = subprocess.run(
            args, 
            capture_output=True, 
            text=True, 
            encoding='utf-8', 
            errors='replace', 
            timeout=timeout,
            check=False, 
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0)
        )
        return (completed.stdout or completed.stderr).strip()
    return safe(run, "Indisponível", "comando do sistema")


def network_info() -> dict[str, Any]:
    data: dict[str, Any] = {"hostname": socket.gethostname(), "ip_local": "Indisponível", "dns": "Indisponível",
                            "gateway": "Indisponível", "internet": "Não verificada"}
    data["ip_local"] = safe(lambda: socket.gethostbyname(socket.gethostname()), data["ip_local"], "IP local")
    if os.name == "nt":
        output = command(["ipconfig", "/all"])
        for line in output.splitlines():
            low, _, value = line.partition(":")
            if "servidores dns" in low.lower() and value.strip(): data["dns"] = value.strip()
            if "gateway padrão" in low.lower() and value.strip(): data["gateway"] = value.strip()
    else:
        data["dns"] = safe(lambda: ", ".join(x.strip() for x in Path("/etc/resolv.conf").read_text(encoding='utf-8', errors='replace').splitlines() if x.startswith("nameserver")), "Indisponível", "DNS")
        data["gateway"] = command(["ip", "route", "show", "default"])
    if requests:
        response = safe(lambda: requests.get("https://www.google.com/generate_204", timeout=5), None, "conectividade externa")
        data["internet"] = f"Disponível (HTTP {response.status_code})" if response is not None else "Sem resposta / indisponível"
    return data


def sipia_audit_info() -> dict[str, Any]:
    """Coleta específica orientada aos quesitos da C.I. 277/2026 e Acesso SIPIA."""
    result: dict[str, Any] = {"url": SIPIA_URL, "http": "Não verificado", "ping": "Não verificado", "status_acesso": "Pendente"}
    if requests:
        response = safe(lambda: requests.get(SIPIA_URL, timeout=10, allow_redirects=True), None, "HTTP SIPIA")
        if response is not None:
            result["http"] = f"HTTP {response.status_code}; {response.url}"
            result["status_acesso"] = "Acessível" if response.status_code < 500 else "Instável/Bloqueado"
        else:
            result["status_acesso"] = "InAcessível / Sem resposta"
    host = "sipiact.mdh.gov.br"
    flag = "-n" if os.name == "nt" else "-c"
    result["ping"] = command(["ping", flag, "1", "-w", "3000" if os.name == "nt" else "3", host], timeout=5)
    return result


def collect() -> dict[str, Any]:
    now = dt.datetime.now().astimezone()
    data: dict[str, Any] = {
        "coletado_em": now.isoformat(timespec="seconds"),
        "contexto_auditoria": "Atendimento C.I. Nº 277/SMAS/GAB/VG/2026 e SIMP nº 016539-006/2025 (Selo UNICEF / SIPIA)",
        "sistema": {
            "hostname": socket.gethostname(), "so": platform.platform(), "versao": platform.version(),
            "python": platform.python_version()
        }
    }
    if psutil:
        ram = psutil.virtual_memory()
        cpu_pct = safe(lambda: psutil.cpu_percent(interval=0.4), 0.0)
        data["sistema"].update({
            "cpu_logica": psutil.cpu_count(),
            "cpu_uso_percentual": cpu_pct,
            "ram_total_bytes": ram.total,
            "ram_usada_percentual": ram.percent,
            "condicao_hardware_apta": ram.total >= 4294967296 and cpu_pct < 90
        })
        disks = []
        for part in safe(psutil.disk_partitions, [], "partições"):
            usage = safe(lambda p=part: psutil.disk_usage(p.mountpoint), None, "uso de disco")
            if usage:
                livre_pct = round(usage.free / usage.total * 100, 1) if usage.total else 0
                disks.append({
                    "dispositivo": part.device,
                    "ponto_montagem": part.mountpoint,
                    "sistema_arquivos": part.fstype,
                    "total_bytes": usage.total,
                    "livre_bytes": usage.free,
                    "livre_percentual": livre_pct,
                    "alerta_espaco": livre_pct < 10
                })
        data["discos"] = disks
    else:
        data["sistema"]["observacao"] = "psutil ausente; dados de hardware limitados."
        data["discos"] = []

    data["rede"] = network_info()
    data["sipia_auditoria"] = sipia_audit_info()
    return data


def gemini(prompt: str, history: list[dict[str, str]] | None = None) -> str:
    key = os.getenv("GEMINI_API_KEY", "").strip()
    if not key:
        return "Parecer de IA indisponível: configure GEMINI_API_KEY no arquivo .env. Dados registrados para laudo técnico manual."
    try:
        from google import genai
        client = genai.Client(api_key=key)
        contents = []
        for item in history or []: contents.append({"role": item["role"], "parts": [{"text": item["text"]}]})
        contents.append({"role": "user", "parts": [{"text": prompt}]})
        response = client.models.generate_content(model="gemini-2.5-flash", contents=contents,
                                                  config={"system_instruction": PERSONA})
        return (response.text or "Não foi possível obter um parecer.").strip()
    except Exception as exc:
        log.exception("Gemini indisponível: %s", exc)
        return "Parecer de IA indisponível no momento. Consulte os dados brutos para preenchimento do relatório ministerial."


# ==========================================
# MÓDULO RAG (Busca Local em Manuais e Docs)
# ==========================================
def buscar_contexto_local(pergunta: str) -> str:
    contextos = []
    if not KNOWLEDGE_DIR.exists():
        return "Nenhuma base de conhecimento local encontrada em docs/."
    
    for arquivo in KNOWLEDGE_DIR.glob("*.*"):
        if arquivo.suffix.lower() in [".txt", ".md"]:
            try:
                conteudo = arquivo.read_text(encoding="utf-8", errors="replace")
                paragrafos = conteudo.split("\n\n")
                for p in paragrafos:
                    palavras_chave = [w.lower() for w in pergunta.split() if len(w) > 3]
                    if any(w in p.lower() for w in palavras_chave):
                        contextos.append(f"[{arquivo.name}]: {p.strip()}")
            except Exception as e:
                log.warning("Erro ao ler o arquivo %s: %s", arquivo.name, e)
                
    return "\n\n".join(contextos[:3]) if contextos else "Nenhum documento correspondente na base local."


def perguntar_com_rag(pergunta: str) -> str:
    contexto_recuperado = buscar_contexto_local(pergunta)
    prompt = f"""Responda à dúvida técnica considerando as exigências da C.I. 277/2026 e do Ministério Público (SIPIA / Selo UNICEF).
CONTEXTO LOCAL:
{contexto_recuperado}

PERGUNTA:
{pergunta}
"""
    return gemini(prompt)


def create_report() -> tuple[Path, Path, Path]:
    data = collect()
    prompt_ia = (
        "Elabore um Parecer Executivo Técnico formal para atender a C.I. Nº 277/SMAS/GAB/VG/2026 e o Procedimento SIMP nº 016539-006/2025 "
        "do Ministério Público (Selo UNICEF). Avalie a adequação dos computadores, estabilidade de rede/internet e acesso ao SIPIA Conselho Tutelar. "
        "Indique claramente se a máquina avaliada está apta ou se requer manutenção/atualização corretiva. "
        "Use apenas estes dados coletados: " + json.dumps(data, ensure_ascii=False)
    )
    opinion = gemini(prompt_ia)
    data["parecer_ia"] = opinion
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    
    txt_path = REPORT_DIR / f"relatorio-auditoria-sipia-{stamp}.txt"
    pdf_path = REPORT_DIR / f"relatorio-auditoria-sipia-{stamp}.pdf"
    docx_path = REPORT_DIR / f"relatorio-auditoria-sipia-{stamp}.docx"
    
    # 1. Salvar TXT
    txt_content = (
        "PREFEITURA MUNICIPAL DE VÁRZEA GRANDE\n"
        "SECRETARIA MUNICIPAL DE ASSISTÊNCIA SOCIAL - SMAS\n"
        "RELATÓRIO TÉCNICO DE AUDITORIA DE TI — ATENDIMENTO C.I. Nº 277/2026 E SIMP 016539-006/2025\n\n"
        "PARECER EXECUTIVO / AUDITORIA SIPIA\n" + opinion + "\n\n"
        "DADOS TÉCNICOS BRUTOS\n" + json.dumps(data, ensure_ascii=False, indent=2)
    )
    safe(lambda: txt_path.write_text(txt_content, encoding="utf-8"), None, "salvar relatório TXT")
    
    # 2. Salvar DOCX (Word)
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("RELATÓRIO DE AUDITORIA TÉCNICA — SIPIA / SELO UNICEF", 0)
        doc.add_paragraph("Prefeitura Municipal de Várzea Grande · Secretaria Municipal de Assistência Social\nReferência: C.I. Nº 277/SMAS/GAB/VG/2026 | SIMP Nº 016539-006/2025\nData da Coleta: " + data["coletado_em"])
        
        doc.add_heading("Parecer Executivo e Adequação Técnica", level=1)
        doc.add_paragraph(opinion)
        
        doc.add_heading("Dados de Infraestrutura Coletados", level=1)
        doc.add_paragraph(json.dumps(data, ensure_ascii=False, indent=2))
        doc.save(str(docx_path))
        log.info("Relatório DOCX de auditoria gerado: %s", docx_path)
    except Exception as exc:
        log.exception("Falha ao gerar DOCX: %s", exc)

    # 3. Salvar PDF
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "B", 14)
        pdf.multi_cell(0, 8, "RELATÓRIO TÉCNICO DE AUDITORIA — SIPIA")
        pdf.set_font("Helvetica", size=9)
        pdf.multi_cell(0, 6, "SMAS Várzea Grande | C.I. 277/2026 | " + data["coletado_em"])
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 7, "Parecer Executivo")
        pdf.set_font("Helvetica", size=9)
        pdf.multi_cell(0, 5, opinion.encode("latin-1", "replace").decode("latin-1"))
        pdf.output(str(pdf_path))
    except Exception as exc:
        log.exception("Falha ao gerar PDF: %s", exc)

    return txt_path, docx_path, pdf_path


def speak(text: str) -> None:
    try:
        import pyttsx3
        engine = pyttsx3.init("sapi5" if os.name == "nt" else None)
        engine.setProperty("rate", 178)
        for voice in engine.getProperty("voices"):
            if "portugu" in (voice.name + voice.id).lower() or "brazil" in (voice.name + voice.id).lower():
                engine.setProperty("voice", voice.id); break
        engine.say(text); engine.runAndWait()
    except Exception as exc: log.exception("Síntese de voz indisponível: %s", exc)


def speech_loop() -> None:
    try:
        import speech_recognition as sr
        try:
            import sounddevice  # noqa: F401
        except Exception as exc: log.warning("sounddevice indisponível: %s", exc)
        recognizer = sr.Recognizer()
        with sr.Microphone() as source:
            recognizer.adjust_for_ambient_noise(source, duration=0.5)
            speak("Jarvis Master pronto para auditoria do SIPIA.")
            while True:
                try:
                    audio = recognizer.listen(source, timeout=4, phrase_time_limit=8)
                    text = recognizer.recognize_google(audio, language="pt-BR")
                    low = text.lower()
                    if "jarvis" in low or "jarbas" in low:
                        if "relatório" in low or "relatorio" in low:
                            paths = create_report(); speak("Relatório de auditoria gerado."); log.info("Relatórios: %s", paths)
                        else: speak("Sistema operacional para atendimento à Circular do SIPIA.")
                except sr.WaitTimeoutError: continue
                except Exception as exc: log.exception("Falha no reconhecimento de voz: %s", exc)
    except Exception as exc:
        log.exception("Captura de microfone indisponível: %s", exc)


def classify_ticket(name: str, description: str) -> str:
    prompt = ("Classifique a prioridade deste chamado como exatamente Alta, Média ou Baixa. Considere impacto, urgência e alcance. "
              "Responda só com prioridade e uma justificativa curta.\nSolicitante: " + name + "\nDescrição: " + description)
    return gemini(prompt)


def helpdesk_poll(stop: threading.Event) -> None:
    endpoint = os.getenv("HELPDESK_URL", "").strip()
    if not endpoint: log.info("Polling Helpdesk desativado."); return
    seen: set[str] = set()
    while not stop.is_set():
        try:
            if not requests: raise RuntimeError("requests indisponível")
            response = requests.get(endpoint, timeout=12); response.raise_for_status(); payload = response.json()
            tickets = payload if isinstance(payload, list) else payload.get("chamados", payload.get("tickets", []))
            for ticket in tickets if isinstance(tickets, list) else []:
                name = str(ticket.get("Nome do Solicitante", ticket.get("nome", ""))).strip()
                desc = str(ticket.get("Descrição Detalhada do Problema", ticket.get("descricao", ""))).strip()
                identity = str(ticket.get("id", "")) or name + "|" + desc
                if name and desc and identity not in seen:
                    seen.add(identity); priority = classify_ticket(name, desc); log.info("Novo chamado (%s): %s", priority, name)
                    speak(f"Novo chamado de {name}. Prioridade {priority}.")
        except Exception as exc: log.exception("Polling Helpdesk falhou: %s", exc)
        stop.wait(30)


def text_chat() -> None:
    history: list[dict[str, str]] = []
    print("Jarvis Master — chat de auditoria SIPIA. Digite sair para encerrar.")
    while True:
        try: prompt = input("Você: ").strip()
        except (EOFError, KeyboardInterrupt): print(); break
        if prompt.lower() in {"sair", "exit", "quit"}: break
        if not prompt: continue
        answer = gemini(prompt, history); print("Jarvis:", answer)
        history.extend([{"role": "user", "text": prompt}, {"role": "model", "text": answer}])


if Flask:
    app = Flask(__name__, static_folder=".", static_url_path="")

    @app.route("/")
    def index():
        return send_from_directory(".", "index.html")

    @app.route("/api/diagnostics")
    def api_diagnostics():
        return jsonify(collect())

    @app.route("/api/relatorio", methods=["POST", "GET"])
    def api_relatorio():
        txt, docx, pdf = create_report()
        return jsonify({
            "mensagem": "Relatórios de auditoria SIPIA gerados com sucesso!",
            "txt": str(txt.name),
            "docx": str(docx.name),
            "pdf": str(pdf.name)
        })

    @app.route("/api/chat-rag", methods=["POST"])
    def api_chat_rag():
        dados = request.get_json() or {}
        pergunta = dados.get("pergunta", "")
        if not pergunta:
            return jsonify({"erro": "Nenhuma pergunta informada."}), 400
        resposta_ia = perguntar_com_rag(pergunta)
        return jsonify({"pergunta": pergunta, "resposta": resposta_ia})

    @app.route("/reports/<path:filename>")
    def baixar_relatorio(filename):
        return send_from_directory(REPORT_DIR, filename, as_attachment=True)

    def run_web_server():
        app.run(host="127.0.0.1", port=5000, debug=False, use_reloader=False)


def main() -> int:
    parser = argparse.ArgumentParser(description="Jarvis Master — apoio técnico e auditoria SIPIA")
    parser.add_argument("--sem-palmas", action="store_true", help="Executa sem captura de voz")
    parser.add_argument("--relatorio", action="store_true", help="Coleta dados de auditoria e gera TXT/DOCX/PDF")
    parser.add_argument("--chat", action="store_true", help="Abre o chat de auditoria")
    parser.add_argument("--web", action="store_true", help="Abre a interface web no navegador")
    args = parser.parse_args()
    try:
        if args.relatorio:
            paths = create_report(); print("Relatórios de auditoria gerados:", *paths); return 0
        if args.chat: text_chat(); return 0

        if Flask:
            threading.Thread(target=run_web_server, name="web-server", daemon=True).start()
            if args.web:
                threading.Timer(1.0, lambda: webbrowser.open("http://127.0.0.1:5000")).start()

        stop = threading.Event()
        poller = threading.Thread(target=helpdesk_poll, args=(stop,), name="helpdesk-poll", daemon=True); poller.start()
        
        if not args.sem_palmas: 
            speech_loop()
        else:
            print("Jarvis Master (Auditoria SIPIA) ativo em http://127.0.0.1:5000. Ctrl+C encerra.")
            while True: time.sleep(1)
    except KeyboardInterrupt: log.info("Encerramento solicitado pelo usuário")
    except Exception as exc: log.exception("Falha geral: %s", exc)
    return 0


if __name__ == "__main__": sys.exit(main())